# -*- coding: utf-8 -*-
"""Carga única del inventario institucional de laboratorios (Fase D).

Lee un Excel con una hoja por laboratorio (columnas PLACA/DESCRIPCIÓN, el
resto se ignora) y crea/actualiza:

- Un `Laboratorio` por hoja (get-or-create por `nombre` == título de la hoja,
  recortado). El Excel NO trae ubicación/capacidad/correo del laboratorio
  -- se usan placeholders (`--ubicacion-default`/`--capacidad-default`/
  `--correo-dominio`) que hay que corregir a mano después desde
  `GestionLaboratoriosScreen`. El Excel es inventario plano, sin ninguna
  noción de Espacio (sub-área).
- Un `Recurso` por fila con PLACA no vacía, dentro del `Laboratorio` de su
  hoja, con un `TipoRecurso` genérico get-or-create ("Equipo
  institucional"). Idempotente por `placa` (columna nueva, único índice
  `uq_recursos_placa`): correr el script dos veces con el mismo Excel NO
  duplica nada, la segunda vez actualiza en vez de insertar.

Los Espacios (sub-áreas) de cada laboratorio (ej. "Estudio de Grabación y
Mezcla 5.1") NO están en el Excel -- se crean a mano después con
`GestionEspaciosScreen`, y los recursos recién importados se les asocian
ahí (Fase A1).

USO (dry-run por defecto, no escribe nada):

    cd backend
    DATABASE_URL=postgresql://postgres:postgres@localhost:5433/reservas_test \
    SECRET_KEY=clave-de-prueba-minimo-32-caracteres \
    SUPABASE_URL=https://x.supabase.co SUPABASE_JWT_SECRET=x-32-caracteres-minimo \
    SUPABASE_SERVICE_ROLE_KEY=x \
    .venv/Scripts/python.exe -m scripts.importar_inventario \
        --archivo "RUTA/A/LABORATORIOS PARQUE I.xlsx" \
        --usuario-id 1

Agregar `--confirmar` para escribir de verdad. `DATABASE_URL` decide el
destino -- NUNCA apuntar esto a `reservas_db` (desarrollo) ni a
producción sin que lo decida explícitamente quien lo corre; verificar el
valor de esa variable antes de agregar `--confirmar`. Este script no
tiene ninguna protección propia contra eso, es responsabilidad de quien
lo ejecuta -- la única barrera es cuál `DATABASE_URL` se exporta.
"""

from __future__ import annotations

import argparse
import sys
from dataclasses import dataclass, field

import openpyxl

sys.path.insert(0, ".")

from app.db import SessionLocal  # noqa: E402
from app.models.laboratorio import Laboratorio  # noqa: E402
from app.models.recurso import Recurso, TipoRecurso  # noqa: E402

TIPO_RECURSO_GENERICO = "Equipo institucional"


@dataclass
class ResumenHoja:
    hoja: str
    laboratorio_nombre: str
    laboratorio_creado: bool = False
    recursos_creados: int = 0
    recursos_actualizados: int = 0
    filas_saltadas: int = 0
    placas_vistas: set[str] = field(default_factory=set)


def _encontrar_encabezado(ws) -> tuple[int, int, int]:
    """Busca la fila de encabezado (la que tiene "PLACA") y devuelve
    (fila, índice_columna_placa, índice_columna_descripcion) -- índices de
    columna 1-based, por POSICIÓN dentro de la fila (`enumerate`), no vía
    `cell.column`: en modo `read_only=True` las celdas vacías son
    `EmptyCell`, que no exponen ese atributo en todas las versiones de
    openpyxl. No asume una fila fija -- confirmado contra el Excel real:
    19 de 20 hojas la tienen en la fila 3, pero una la tiene en la fila 4.
    """
    for idx_fila, fila in enumerate(ws.iter_rows(min_row=1, max_row=10), start=1):
        celdas = {
            idx_col: (str(c.value).strip().upper() if c.value else "")
            for idx_col, c in enumerate(fila, start=1)
        }
        col_placa = next((col for col, val in celdas.items() if val == "PLACA"), None)
        if col_placa is None:
            continue
        col_desc = next(
            (col for col, val in celdas.items() if val.startswith("DESCRIPCI")), None
        )
        if col_desc is None:
            raise ValueError(f"Hoja {ws.title!r}: fila con PLACA pero sin columna DESCRIPCIÓN")
        return idx_fila, col_placa, col_desc
    raise ValueError(f"Hoja {ws.title!r}: no se encontró una fila con encabezado PLACA en las primeras 10 filas")


def _get_or_create_tipo_recurso(db) -> TipoRecurso:
    tipo = db.query(TipoRecurso).filter(TipoRecurso.nombre == TIPO_RECURSO_GENERICO).first()
    if tipo is not None:
        return tipo
    tipo = TipoRecurso(nombre=TIPO_RECURSO_GENERICO, descripcion="Creado por scripts/importar_inventario.py", activo="activo")
    db.add(tipo)
    db.flush()
    return tipo


def _get_or_create_laboratorio(db, *, nombre: str, ubicacion: str, capacidad: int, correo_dominio: str, resumen: ResumenHoja) -> Laboratorio:
    laboratorio = db.query(Laboratorio).filter(Laboratorio.nombre == nombre).first()
    if laboratorio is not None:
        return laboratorio
    slug = "".join(c if c.isalnum() else "-" for c in nombre.lower()).strip("-")
    laboratorio = Laboratorio(
        nombre=nombre,
        ubicacion=ubicacion,
        capacidad=capacidad,
        estado="activo",
        correo=f"{slug}@{correo_dominio}",
        # Horario por defecto del modelo (lunes-sábado 7-19h, ver
        # app/models/laboratorio.py) -- se corrige a mano si un laboratorio
        # real atiende distinto.
    )
    db.add(laboratorio)
    db.flush()
    resumen.laboratorio_creado = True
    return laboratorio


def _importar_hoja(db, ws, *, usuario_id: int, ubicacion_default: str, capacidad_default: int, correo_dominio: str, tipo_recurso: TipoRecurso) -> ResumenHoja:
    nombre_laboratorio = ws.title.strip()
    resumen = ResumenHoja(hoja=ws.title, laboratorio_nombre=nombre_laboratorio)
    fila_encabezado, col_placa, col_desc = _encontrar_encabezado(ws)

    laboratorio = _get_or_create_laboratorio(
        db,
        nombre=nombre_laboratorio,
        ubicacion=ubicacion_default,
        capacidad=capacidad_default,
        correo_dominio=correo_dominio,
        resumen=resumen,
    )

    for fila in ws.iter_rows(min_row=fila_encabezado + 1):
        placa_val = fila[col_placa - 1].value
        desc_val = fila[col_desc - 1].value
        placa = str(placa_val).strip() if placa_val is not None else ""
        descripcion = str(desc_val).strip() if desc_val is not None else ""
        if not placa or not descripcion:
            resumen.filas_saltadas += 1
            continue
        if placa in resumen.placas_vistas:
            # No debería pasar (confirmado sin duplicados en el Excel real),
            # pero si un Excel futuro sí los tiene, no duplicar el conteo.
            resumen.filas_saltadas += 1
            continue
        resumen.placas_vistas.add(placa)

        recurso = db.query(Recurso).filter(Recurso.placa == placa).first()
        if recurso is None:
            recurso = Recurso(
                nombre=descripcion[:100],
                laboratorio_id=laboratorio.id,
                tipo_recurso_id=tipo_recurso.id,
                capacidad=1,
                estado="activo",
                placa=placa,
                created_by=usuario_id,
                update_by=usuario_id,
            )
            db.add(recurso)
            resumen.recursos_creados += 1
        else:
            recurso.nombre = descripcion[:100]
            recurso.laboratorio_id = laboratorio.id
            recurso.update_by = usuario_id
            resumen.recursos_actualizados += 1

    db.flush()
    return resumen


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--archivo", required=True, help="Ruta al .xlsx")
    parser.add_argument("--usuario-id", required=True, type=int, help="Usuario admin al que se atribuye la carga (created_by/update_by)")
    parser.add_argument("--confirmar", action="store_true", help="Escribe de verdad. Sin esto, solo imprime el resumen (dry-run)")
    parser.add_argument("--capacidad-default", type=int, default=15, help="Capacidad placeholder para laboratorios nuevos (default: 15)")
    parser.add_argument("--ubicacion-default", default="Parque i — piso por definir", help="Ubicación placeholder para laboratorios nuevos")
    parser.add_argument("--correo-dominio", default="pendiente.itm.edu.co", help="Dominio placeholder para el correo de laboratorios nuevos")
    args = parser.parse_args()

    wb = openpyxl.load_workbook(args.archivo, data_only=True, read_only=True)

    db = SessionLocal()
    try:
        tipo_recurso = _get_or_create_tipo_recurso(db)
        resumenes = [
            _importar_hoja(
                db,
                ws,
                usuario_id=args.usuario_id,
                ubicacion_default=args.ubicacion_default,
                capacidad_default=args.capacidad_default,
                correo_dominio=args.correo_dominio,
                tipo_recurso=tipo_recurso,
            )
            for ws in wb.worksheets
        ]

        total_creados = sum(r.recursos_creados for r in resumenes)
        total_actualizados = sum(r.recursos_actualizados for r in resumenes)
        total_saltados = sum(r.filas_saltadas for r in resumenes)
        laboratorios_nuevos = sum(1 for r in resumenes if r.laboratorio_creado)

        print(f"{'CONFIRMAR' if args.confirmar else 'DRY-RUN'} -- {len(resumenes)} hojas procesadas\n")
        for r in resumenes:
            marca_laboratorio = "NUEVO" if r.laboratorio_creado else "ya existía"
            print(
                f"  {r.hoja!r} -> Laboratorio {r.laboratorio_nombre!r} ({marca_laboratorio}): "
                f"{r.recursos_creados} creados, {r.recursos_actualizados} actualizados, "
                f"{r.filas_saltadas} filas saltadas"
            )
        print(
            f"\nTOTAL: {laboratorios_nuevos} laboratorios nuevos, {total_creados} recursos creados, "
            f"{total_actualizados} recursos actualizados, {total_saltados} filas saltadas"
        )

        if args.confirmar:
            db.commit()
            print("\nEscrito en la base de datos.")
        else:
            db.rollback()
            print("\nDry-run: nada se escribió. Agregar --confirmar para aplicar los cambios.")
    finally:
        db.close()


if __name__ == "__main__":
    main()
