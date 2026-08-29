# -*- coding: utf-8 -*-
"""Pruebas de `GET /reservas/mis-reservas/export` (2026-08-29) -- mismo
patrón de verificación que `test_admin_dashboard_export.py`: parsear el
archivo real de vuelta, no solo comprobar el status.
"""

import csv
import io

from openpyxl import load_workbook

from tests.conftest import cookies_para, crear_espacio, crear_recurso, crear_usuario, fecha_habilitada, payload_reserva


def test_export_csv_incluye_la_reserva_propia(client, db):
    espacio = crear_espacio(db, nombre="Espacio Export Reservas")
    usuario = crear_usuario(db, username="export_reservas_user", email="export_reservas_user@example.com", rol="usuario")
    recurso = crear_recurso(db, espacio=espacio, usuario=usuario, nombre="Recurso Export Reservas")
    headers = cookies_para(usuario)
    client.post("/reservas", json=payload_reserva(recurso.id, fecha_habilitada()), headers=headers)

    respuesta = client.get("/reservas/mis-reservas/export", params={"formato": "csv"}, headers=headers)

    assert respuesta.status_code == 200
    assert respuesta.headers["content-type"].startswith("text/csv")
    filas = list(csv.reader(io.StringIO(respuesta.content.decode("utf-8-sig"))))
    assert filas[0] == ["Fecha", "Hora inicio", "Hora fin", "Espacio", "Recursos", "Estado", "Asistentes"]
    assert any(fila[3] == espacio.nombre for fila in filas[1:])


def test_export_xlsx_devuelve_libro_valido(client, db):
    espacio = crear_espacio(db, nombre="Espacio Export XLSX")
    usuario = crear_usuario(db, username="export_reservas_xlsx", email="export_reservas_xlsx@example.com", rol="usuario")
    recurso = crear_recurso(db, espacio=espacio, usuario=usuario, nombre="Recurso Export XLSX")
    headers = cookies_para(usuario)
    client.post("/reservas", json=payload_reserva(recurso.id, fecha_habilitada()), headers=headers)

    respuesta = client.get("/reservas/mis-reservas/export", params={"formato": "xlsx"}, headers=headers)

    assert respuesta.status_code == 200
    libro = load_workbook(io.BytesIO(respuesta.content))
    assert "Mis reservas" in libro.sheetnames
    hoja = libro["Mis reservas"]
    assert hoja.cell(row=1, column=1).value == "Fecha"


def test_export_no_incluye_reservas_de_otro(client, db):
    espacio = crear_espacio(db, nombre="Espacio Export Ajeno")
    dueno = crear_usuario(db, username="export_dueno", email="export_dueno@example.com", rol="usuario")
    otro = crear_usuario(db, username="export_otro", email="export_otro@example.com", rol="usuario")
    recurso = crear_recurso(db, espacio=espacio, usuario=dueno, nombre="Recurso Export Ajeno")
    client.post("/reservas", json=payload_reserva(recurso.id, fecha_habilitada()), headers=cookies_para(dueno))

    respuesta = client.get("/reservas/mis-reservas/export", params={"formato": "csv"}, headers=cookies_para(otro))

    filas = list(csv.reader(io.StringIO(respuesta.content.decode("utf-8-sig"))))
    assert len(filas) == 1  # solo el encabezado


def test_export_sin_sesion_da_401(client, db):
    respuesta = client.get("/reservas/mis-reservas/export", params={"formato": "csv"})
    assert respuesta.status_code == 401
