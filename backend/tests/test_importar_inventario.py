# -*- coding: utf-8 -*-
"""Pruebas de `scripts/importar_inventario.py` (Fase D).

No se prueba `main()` completo (arma su propio `SessionLocal()`, fuera del
aislamiento por `db`/`TRUNCATE` de `conftest.py`) -- se prueban las
funciones internas directamente contra la sesión de la fixture `db`, con
workbooks sintéticos minúsculos guardados en `tmp_path` y reabiertos con
`read_only=True` (mismo modo que usa `main()` contra el Excel real, para
cubrir de verdad el bug de `EmptyCell` ya corregido en
`_encontrar_encabezado`).
"""

import openpyxl
import pytest

from app.models.espacio import Espacio
from app.models.recurso import Recurso
from scripts.importar_inventario import (
    _encontrar_encabezado,
    _get_or_create_espacio,
    _get_or_create_tipo_recurso,
    _importar_hoja,
    ResumenHoja,
)
from tests.conftest import crear_usuario


def _hoja_desde_filas(tmp_path, nombre_hoja, filas, nombre_archivo="inventario.xlsx"):
    """Crea un .xlsx con una única hoja `nombre_hoja` con `filas` (lista de
    listas), lo guarda en `tmp_path` y devuelve la hoja reabierta en modo
    `read_only=True` (igual que `main()`)."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = nombre_hoja
    for fila in filas:
        ws.append(fila)
    ruta = tmp_path / nombre_archivo
    wb.save(ruta)

    wb_reabierto = openpyxl.load_workbook(ruta, data_only=True, read_only=True)
    return wb_reabierto[nombre_hoja]


class TestEncontrarEncabezado:
    def test_encabezado_en_fila_tipica(self, tmp_path):
        ws = _hoja_desde_filas(
            tmp_path,
            "LAB PRUEBA",
            [
                ["LABORATORIO PARQUE I"],
                [],
                ["PLACA", "DESCRIPCIÓN", "BODEGA"],
                ["001", "SILLA"],
            ],
        )
        fila, col_placa, col_desc = _encontrar_encabezado(ws)
        assert fila == 3
        assert col_placa == 1
        assert col_desc == 2

    def test_encabezado_desplazado_no_hardcodeado(self, tmp_path):
        """Confirmado contra el Excel real: 19/20 hojas tienen el
        encabezado en la fila 3, pero una lo tiene en la fila 4 -- este
        test cubre justo ese caso, con una fila en blanco extra antes."""
        ws = _hoja_desde_filas(
            tmp_path,
            "LAB DESPLAZADO",
            [
                ["LABORATORIO PARQUE I"],
                [],
                [],
                ["PLACA", "DESCRIPCIÓN"],
                ["002", "MESA"],
            ],
        )
        fila, col_placa, col_desc = _encontrar_encabezado(ws)
        assert fila == 4

    def test_columnas_no_adyacentes_por_celdas_vacias(self, tmp_path):
        """Fuerza celdas vacías (`EmptyCell` en modo read_only) entre
        columnas reales -- es exactamente el escenario que rompía la
        versión anterior de `_encontrar_encabezado` (`c.column` no existe
        en `EmptyCell`)."""
        ws = _hoja_desde_filas(
            tmp_path,
            "LAB HUECOS",
            [
                ["PLACA", None, "DESCRIPCIÓN", None, "COSTO"],
                ["003", None, "PROYECTOR", None, "1000"],
            ],
        )
        fila, col_placa, col_desc = _encontrar_encabezado(ws)
        assert fila == 1
        assert col_placa == 1
        assert col_desc == 3

    def test_sin_placa_lanza_error(self, tmp_path):
        ws = _hoja_desde_filas(tmp_path, "LAB SIN PLACA", [["FOO", "BAR"], ["1", "2"]])
        with pytest.raises(ValueError, match="no se encontró"):
            _encontrar_encabezado(ws)

    def test_placa_sin_descripcion_lanza_error(self, tmp_path):
        ws = _hoja_desde_filas(tmp_path, "LAB SIN DESC", [["PLACA", "BODEGA"], ["1", "A"]])
        with pytest.raises(ValueError, match="sin columna DESCRIPCIÓN"):
            _encontrar_encabezado(ws)


class TestGetOrCreateEspacio:
    def test_crea_espacio_nuevo_con_placeholders(self, db):
        resumen = ResumenHoja(hoja="X", espacio_nombre="LAB NUEVO")
        espacio = _get_or_create_espacio(
            db,
            nombre="LAB NUEVO",
            ubicacion="Parque i — piso por definir",
            capacidad=15,
            correo_dominio="pendiente.itm.edu.co",
            resumen=resumen,
        )
        assert espacio.id is not None
        assert espacio.ubicacion == "Parque i — piso por definir"
        assert espacio.capacidad == 15
        assert espacio.correo == "lab-nuevo@pendiente.itm.edu.co"
        assert espacio.estado == "activo"
        assert resumen.espacio_creado is True

    def test_reusa_espacio_existente_por_nombre(self, db):
        db.add(
            Espacio(
                nombre="LAB YA EXISTE",
                ubicacion="Ubicación real",
                capacidad=40,
                estado="activo",
                modalidad_reserva="equipos",
                correo="labyaexiste@itm.edu.co",
            )
        )
        db.commit()

        resumen = ResumenHoja(hoja="X", espacio_nombre="LAB YA EXISTE")
        espacio = _get_or_create_espacio(
            db,
            nombre="LAB YA EXISTE",
            ubicacion="Parque i — piso por definir",
            capacidad=15,
            correo_dominio="pendiente.itm.edu.co",
            resumen=resumen,
        )
        assert espacio.ubicacion == "Ubicación real"
        assert espacio.capacidad == 40
        assert resumen.espacio_creado is False
        assert db.query(Espacio).filter(Espacio.nombre == "LAB YA EXISTE").count() == 1


class TestImportarHoja:
    def test_crea_espacio_y_recursos_salta_filas_vacias(self, db, tmp_path):
        ws = _hoja_desde_filas(
            tmp_path,
            "LAB IMPORT",
            [
                ["PLACA", "DESCRIPCIÓN"],
                ["100", "PROYECTOR EPSON"],
                ["101", "PARLANTE JBL"],
                ["", "SIN PLACA -- SE SALTA"],
                ["102", ""],
                ["103", "MICRÓFONO SM 57"],
            ],
        )
        tipo = _get_or_create_tipo_recurso(db)
        usuario = crear_usuario(db, username="importador", email="importador@itm.edu.co", rol="admin")

        resumen = _importar_hoja(
            db,
            ws,
            usuario_id=usuario.id,
            ubicacion_default="Parque i — piso por definir",
            capacidad_default=15,
            correo_dominio="pendiente.itm.edu.co",
            tipo_recurso=tipo,
        )

        assert resumen.espacio_creado is True
        assert resumen.recursos_creados == 3
        assert resumen.recursos_actualizados == 0
        assert resumen.filas_saltadas == 2

        espacio = db.query(Espacio).filter(Espacio.nombre == "LAB IMPORT").one()
        recursos = db.query(Recurso).filter(Recurso.espacio_id == espacio.id).all()
        assert {r.placa for r in recursos} == {"100", "101", "103"}
        assert all(r.tipo_recurso_id == tipo.id for r in recursos)
        assert all(r.capacidad == 1 and r.estado == "activo" for r in recursos)

    def test_placa_duplicada_dentro_de_la_misma_hoja_se_saltea(self, db, tmp_path):
        ws = _hoja_desde_filas(
            tmp_path,
            "LAB DUPLICADO",
            [
                ["PLACA", "DESCRIPCIÓN"],
                ["200", "SILLA"],
                ["200", "SILLA DUPLICADA"],
            ],
        )
        tipo = _get_or_create_tipo_recurso(db)
        usuario = crear_usuario(db, username="importador", email="importador@itm.edu.co", rol="admin")

        resumen = _importar_hoja(
            db,
            ws,
            usuario_id=usuario.id,
            ubicacion_default="Parque i — piso por definir",
            capacidad_default=15,
            correo_dominio="pendiente.itm.edu.co",
            tipo_recurso=tipo,
        )

        assert resumen.recursos_creados == 1
        assert resumen.filas_saltadas == 1

    def test_reimportar_mismo_excel_es_idempotente_por_placa(self, db, tmp_path):
        """Correr el import dos veces con el mismo Excel no duplica
        recursos -- la segunda vuelta actualiza en vez de insertar."""
        ws1 = _hoja_desde_filas(
            tmp_path,
            "LAB REPETIDO",
            [["PLACA", "DESCRIPCIÓN"], ["300", "TABLET SAMSUNG"]],
            nombre_archivo="v1.xlsx",
        )
        tipo = _get_or_create_tipo_recurso(db)
        usuario = crear_usuario(db, username="importador", email="importador@itm.edu.co", rol="admin")
        primero = _importar_hoja(
            db, ws1, usuario_id=usuario.id, ubicacion_default="X", capacidad_default=15,
            correo_dominio="pendiente.itm.edu.co", tipo_recurso=tipo,
        )
        db.commit()
        assert primero.recursos_creados == 1
        assert primero.recursos_actualizados == 0

        ws2 = _hoja_desde_filas(
            tmp_path,
            "LAB REPETIDO",
            [["PLACA", "DESCRIPCIÓN"], ["300", "TABLET SAMSUNG (renombrada)"]],
            nombre_archivo="v2.xlsx",
        )
        segundo = _importar_hoja(
            db, ws2, usuario_id=usuario.id, ubicacion_default="X", capacidad_default=15,
            correo_dominio="pendiente.itm.edu.co", tipo_recurso=tipo,
        )
        db.commit()
        assert segundo.recursos_creados == 0
        assert segundo.recursos_actualizados == 1
        assert segundo.espacio_creado is False

        recursos = db.query(Recurso).filter(Recurso.placa == "300").all()
        assert len(recursos) == 1
        assert recursos[0].nombre == "TABLET SAMSUNG (renombrada)"

    def test_dry_run_no_persiste_nada(self, db, tmp_path):
        """Réplica del contrato dry-run de `main()`: sin `--confirmar`, se
        hace `rollback()` en vez de `commit()` y no debe quedar nada."""
        ws = _hoja_desde_filas(
            tmp_path,
            "LAB DRY RUN",
            [["PLACA", "DESCRIPCIÓN"], ["400", "CÁMARA"]],
        )
        tipo = _get_or_create_tipo_recurso(db)
        usuario = crear_usuario(db, username="importador", email="importador@itm.edu.co", rol="admin")
        _importar_hoja(
            db, ws, usuario_id=usuario.id, ubicacion_default="X", capacidad_default=15,
            correo_dominio="pendiente.itm.edu.co", tipo_recurso=tipo,
        )
        db.rollback()

        assert db.query(Espacio).filter(Espacio.nombre == "LAB DRY RUN").count() == 0
        assert db.query(Recurso).filter(Recurso.placa == "400").count() == 0
