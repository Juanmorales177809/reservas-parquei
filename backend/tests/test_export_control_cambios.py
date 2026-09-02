# -*- coding: utf-8 -*-
"""Pruebas de `GET /admin/control-cambios/export` (2026-08-29)."""

import csv
import io

from openpyxl import load_workbook

from tests.conftest import cookies_para, crear_laboratorio, crear_recurso, crear_usuario, fecha_habilitada, payload_reserva


def _admin_headers(db):
    admin = crear_usuario(db, username="export_auditoria_admin", email="export_auditoria_admin@example.com", rol="admin")
    return admin, cookies_para(admin)


def test_export_csv_incluye_cambios_reales(client, db):
    admin, headers = _admin_headers(db)
    laboratorio = crear_laboratorio(db, nombre="Laboratorio Auditoria Export")
    recurso = crear_recurso(db, laboratorio=laboratorio, usuario=admin, nombre="Recurso Auditoria Export")
    client.post("/reservas", json=payload_reserva(recurso.id, fecha_habilitada()), headers=headers)

    respuesta = client.get("/admin/control-cambios/export", params={"formato": "csv"}, headers=headers)

    assert respuesta.status_code == 200
    filas = list(csv.reader(io.StringIO(respuesta.content.decode("utf-8-sig"))))
    assert filas[0] == ["Fecha", "Actor", "Acción", "Entidad", "ID entidad", "Descripción"]
    assert any(fila[2] == "crear" and fila[3] == "reserva" for fila in filas[1:])


def test_export_xlsx_devuelve_libro_valido(client, db):
    admin, headers = _admin_headers(db)

    respuesta = client.get("/admin/control-cambios/export", params={"formato": "xlsx"}, headers=headers)

    assert respuesta.status_code == 200
    libro = load_workbook(io.BytesIO(respuesta.content))
    assert "Auditoría" in libro.sheetnames


def test_export_usuario_comun_da_403(client, db):
    usuario = crear_usuario(db, username="export_auditoria_usuario", email="export_auditoria_usuario@example.com", rol="usuario")
    respuesta = client.get("/admin/control-cambios/export", params={"formato": "csv"}, headers=cookies_para(usuario))
    assert respuesta.status_code == 403
