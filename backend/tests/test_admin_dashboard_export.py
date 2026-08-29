# -*- coding: utf-8 -*-
"""Pruebas del export del dashboard admin/gestor a CSV/Excel (2026-08-29).

Usa la misma fuente de datos que `/summary` (`_construir_resumen`); estas
pruebas verifican el transporte (status, content-type, filename, formato de
archivo real) sin duplicar la lógica de agregación ya cubierta en
`test_admin_dashboard_ocupacion.py`.
"""

import csv
import io

from openpyxl import load_workbook

from tests.conftest import cookies_para, crear_espacio, crear_recurso, crear_usuario


def _admin_y_headers(db):
    admin = crear_usuario(db, username="admin_export", email="admin_export@example.com", rol="admin")
    return admin, cookies_para(admin)


def _gestor_con_espacio(db):
    espacio = crear_espacio(db, nombre="Espacio Export")
    gestor = crear_usuario(
        db, username="gestor_export", email="gestor_export@example.com", rol="gestor", espacio_id=espacio.id
    )
    return gestor, espacio, cookies_para(gestor)


def test_export_admin_csv_devuelve_archivo_csv_valido(client, db):
    admin, headers = _admin_y_headers(db)
    espacio = crear_espacio(db, nombre="Espacio CSV")
    crear_recurso(db, espacio=espacio, usuario=admin, nombre="Recurso CSV")

    respuesta = client.get("/admin/dashboard/export", params={"formato": "csv"}, headers=headers)

    assert respuesta.status_code == 200
    assert respuesta.headers["content-type"].startswith("text/csv")
    assert "attachment; filename=dashboard_" in respuesta.headers["content-disposition"]
    filas = list(csv.reader(io.StringIO(respuesta.content.decode("utf-8-sig"))))
    assert filas[0] == ["Métrica", "Valor"]
    assert any(fila[:1] == ["Recursos activos"] for fila in filas)


def test_export_admin_xlsx_devuelve_libro_valido_con_hojas(client, db):
    admin, headers = _admin_y_headers(db)

    respuesta = client.get("/admin/dashboard/export", params={"formato": "xlsx"}, headers=headers)

    assert respuesta.status_code == 200
    assert respuesta.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    libro = load_workbook(io.BytesIO(respuesta.content))
    assert "Resumen" in libro.sheetnames
    assert "Ocupación por día y hora" in libro.sheetnames


def test_export_gestor_scoped_a_su_espacio(client, db):
    gestor, espacio, headers = _gestor_con_espacio(db)

    respuesta = client.get("/gestion/dashboard/export", params={"formato": "csv"}, headers=headers)

    assert respuesta.status_code == 200
    filas = list(csv.reader(io.StringIO(respuesta.content.decode("utf-8-sig"))))
    assert any(fila == ["Espacio", espacio.nombre] for fila in filas)


def test_export_usuario_comun_da_403(client, db):
    usuario = crear_usuario(db, username="usuario_export", email="usuario_export@example.com", rol="usuario")
    headers = cookies_para(usuario)

    respuesta = client.get("/admin/dashboard/export", params={"formato": "csv"}, headers=headers)

    assert respuesta.status_code == 403


def test_export_sin_formato_da_422(client, db):
    admin, headers = _admin_y_headers(db)

    respuesta = client.get("/admin/dashboard/export", headers=headers)

    assert respuesta.status_code == 422
