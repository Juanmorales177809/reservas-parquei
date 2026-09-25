"""Lógica de importaciones masivas (API-12, contrato §4).

Administration valida el archivo, orquesta la carga, presenta el resumen y
conserva la trazabilidad (`RN-IMP-10`); no define las reglas propias de
cada catálogo. Las entidades se escriben siempre en el módulo propietario
—`researchs` para proyectos/semilleros, `resources` para equipos— con sus
propias funciones de creación, nunca con SQL directo que eluda sus reglas.

La carga es de dos pasos (validar → confirmar) sin que la confirmación
reciba cuerpo. `administration.importacion_resultados.datos` guarda los
valores normalizados de cada fila que vaya a escribir algo, precisamente
para que confirmar no necesite el archivo otra vez (ver
`backend/migrations/012_importacion_resultados_datos.sql`).
"""

from __future__ import annotations

import io
from datetime import date, datetime, timezone

import openpyxl

from app.core.deps import ContextoAutenticado
from app.core.errors import Conflicto, NoEncontrado, SolicitudInvalida, Validacion
from app.modules.administration import repository as repo
from app.modules.researchs import repository as inv_repo
from app.modules.resources import repository as rec_repo

CATALOGOS = ("PROYECTOS", "SEMILLEROS", "EQUIPOS")

_COLUMNAS_INVESTIGACION = ("codigo", "nombre", "estado")
_COLUMNAS_EQUIPOS = ("PLACA", "DESCRIPCIÓN", "CODIGO BODEGA", "CENTRO DE COSTOS", "FECHA INICIO", "COSTO")


def _ahora() -> datetime:
    return datetime.now(timezone.utc)


# --- Lectura del archivo -----------------------------------------------------------


def _leer_hoja(contenido: bytes):
    try:
        libro = openpyxl.load_workbook(io.BytesIO(contenido), data_only=True, read_only=True)
    except Exception as exc:
        raise Validacion("El archivo no es legible.") from exc
    return libro.active


def _encabezados(hoja) -> dict[str, int]:
    """`{encabezado_normalizado: índice_de_columna}` de la primera fila."""
    fila = next(hoja.iter_rows(min_row=1, max_row=1), None)
    if fila is None:
        raise Validacion("El archivo no tiene encabezados.")
    encabezados = {}
    for i, celda in enumerate(fila):
        if celda.value is not None:
            encabezados[str(celda.value).strip().lower()] = i
    return encabezados


def _filas_de_datos(hoja):
    for fila in hoja.iter_rows(min_row=2):
        if all(c.value is None for c in fila):
            continue
        yield fila


# --- §4.1 Validar --------------------------------------------------------------------


def _validar_columnas(encabezados: dict, esperadas: tuple[str, ...]) -> None:
    presentes = set(encabezados.keys())
    faltantes = [c for c in esperadas if c.lower() not in presentes]
    if faltantes:
        raise Validacion(f"Faltan columnas: {', '.join(faltantes)}.")


def _valor(fila, encabezados: dict, columna: str):
    idx = encabezados.get(columna.lower())
    if idx is None or idx >= len(fila):
        return None
    return fila[idx].value


def _texto(valor) -> str:
    return str(valor).strip() if valor is not None else ""


def _procesar_investigacion(db, hoja, encabezados: dict, catalogo: str) -> list[dict]:
    obtener_por_codigo = inv_repo.obtener_proyecto_por_codigo if catalogo == "PROYECTOS" else inv_repo.obtener_semillero_por_codigo
    codigos_vistos: dict[str, int] = {}
    resultados = []
    for numero_fila, fila in enumerate(_filas_de_datos(hoja), start=2):
        codigo_raw = _texto(_valor(fila, encabezados, "codigo"))
        nombre = _texto(_valor(fila, encabezados, "nombre"))
        estado_raw = _texto(_valor(fila, encabezados, "estado")).upper()

        codigo = codigo_raw.strip().upper()
        if not codigo:
            resultados.append({"numero_fila": numero_fila, "codigo": None, "resultado": "ERROR", "detalle": "La fila no trae código.", "datos": None})
            continue
        if not nombre:
            resultados.append({"numero_fila": numero_fila, "codigo": codigo, "resultado": "ERROR", "detalle": "La fila no trae nombre.", "datos": None})
            continue
        if estado_raw not in ("ACTIVO", "INACTIVO"):
            resultados.append({"numero_fila": numero_fila, "codigo": codigo, "resultado": "ERROR", "detalle": "El estado debe ser ACTIVO o INACTIVO.", "datos": None})
            continue
        if codigo in codigos_vistos:
            resultados.append({"numero_fila": numero_fila, "codigo": codigo, "resultado": "ERROR", "detalle": f"Código duplicado con la fila {codigos_vistos[codigo]}.", "datos": None})
            continue
        codigos_vistos[codigo] = numero_fila

        estado = estado_raw == "ACTIVO"
        existente = obtener_por_codigo(db, codigo)
        datos = {"codigo": codigo, "nombre": nombre, "estado": estado}
        if existente is None:
            resultados.append({"numero_fila": numero_fila, "codigo": codigo, "resultado": "CREADO", "detalle": None, "datos": datos})
        elif existente.estado and not estado:
            resultados.append({"numero_fila": numero_fila, "codigo": codigo, "resultado": "DESACTIVADO", "detalle": None, "datos": datos})
        else:
            resultados.append({"numero_fila": numero_fila, "codigo": codigo, "resultado": "ACTUALIZADO", "detalle": None, "datos": datos})
    return resultados


def _fecha(valor) -> str | None:
    if valor is None or _texto(valor) == "":
        return None
    if isinstance(valor, datetime):
        return valor.date().isoformat()
    if isinstance(valor, date):
        return valor.isoformat()
    texto = _texto(valor)
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
        try:
            return datetime.strptime(texto, fmt).date().isoformat()
        except ValueError:
            continue
    return "__invalida__"


def _procesar_equipos(db, hoja, encabezados: dict) -> list[dict]:
    placas_vistas: dict[str, int] = {}
    resultados = []
    for numero_fila, fila in enumerate(_filas_de_datos(hoja), start=2):
        placa = _texto(_valor(fila, encabezados, "PLACA"))
        descripcion = _texto(_valor(fila, encabezados, "DESCRIPCIÓN"))
        bodega = _texto(_valor(fila, encabezados, "CODIGO BODEGA")) or None
        centro_costo = _texto(_valor(fila, encabezados, "CENTRO DE COSTOS")) or None
        fecha_compra = _fecha(_valor(fila, encabezados, "FECHA INICIO"))
        # COSTO se lee y se descarta (RN-IMP-11 de administration).

        if not placa:
            resultados.append({"numero_fila": numero_fila, "codigo": None, "resultado": "ERROR", "detalle": "La fila no trae placa.", "datos": None})
            continue
        if not descripcion:
            resultados.append({"numero_fila": numero_fila, "codigo": placa, "resultado": "ERROR", "detalle": "La fila no trae descripción.", "datos": None})
            continue
        if fecha_compra == "__invalida__":
            resultados.append({"numero_fila": numero_fila, "codigo": placa, "resultado": "ERROR", "detalle": "FECHA INICIO no es una fecha válida.", "datos": None})
            continue
        if placa in placas_vistas:
            resultados.append({"numero_fila": numero_fila, "codigo": placa, "resultado": "ERROR", "detalle": f"Placa duplicada con la fila {placas_vistas[placa]}.", "datos": None})
            continue
        placas_vistas[placa] = numero_fila

        existente = rec_repo.obtener_equipo_por_placa(db, placa)
        datos = {
            "placa": placa, "nombre_equipo": descripcion, "bodega": bodega,
            "centro_costo": centro_costo, "fecha_compra": fecha_compra,
        }
        resultado = "ACTUALIZADO" if existente is not None else "CREADO"
        resultados.append({"numero_fila": numero_fila, "codigo": placa, "resultado": resultado, "detalle": None, "datos": datos})
    return resultados


def validar_importacion(
    db, *, catalogo: str, id_unidad: int | None, archivo_nombre: str, archivo_bytes: bytes, contexto: ContextoAutenticado,
) -> dict:
    if catalogo not in CATALOGOS:
        raise Validacion(f"'catalogo' debe ser uno de {', '.join(CATALOGOS)}.")
    if catalogo == "EQUIPOS":
        if id_unidad is None:
            raise SolicitudInvalida("id_unidad es obligatorio para el catálogo EQUIPOS.")
        if rec_repo.obtener_unidad(db, id_unidad) is None:
            raise NoEncontrado("La unidad no existe.")
    else:
        id_unidad = None

    hoja = _leer_hoja(archivo_bytes)
    encabezados = _encabezados(hoja)
    if catalogo == "EQUIPOS":
        _validar_columnas(encabezados, _COLUMNAS_EQUIPOS)
        resultados = _procesar_equipos(db, hoja, encabezados)
        if id_unidad is not None:
            for r in resultados:
                if r["datos"] is not None:
                    r["datos"]["id_unidad"] = id_unidad
    else:
        _validar_columnas(encabezados, _COLUMNAS_INVESTIGACION)
        resultados = _procesar_investigacion(db, hoja, encabezados, catalogo)

    importacion = repo.crear_importacion(db, actor_cuenta_id=contexto.id_cuenta, catalogo=catalogo, archivo_referencia=archivo_nombre)
    for r in resultados:
        repo.agregar_resultado_fila(db, importacion_id=importacion.id, **r)
    db.commit()

    a_crear = sum(1 for r in resultados if r["resultado"] == "CREADO")
    a_actualizar = sum(1 for r in resultados if r["resultado"] in ("ACTUALIZADO", "DESACTIVADO"))
    con_error = sum(1 for r in resultados if r["resultado"] == "ERROR")
    return {
        "id": importacion.id, "catalogo": catalogo, "id_unidad": id_unidad,
        "archivo_referencia": archivo_nombre, "confirmable": con_error == 0,
        "totales": {"a_crear": a_crear, "a_actualizar": a_actualizar, "con_error": con_error},
        "resultados": [{"numero_fila": r["numero_fila"], "codigo": r["codigo"], "resultado": r["resultado"], "detalle": r["detalle"]} for r in resultados],
    }


# --- §4.2 Confirmar --------------------------------------------------------------------


def _confirmar_investigacion(db, catalogo: str, resultados) -> tuple[int, int, int]:
    crear = inv_repo.crear_proyecto if catalogo == "PROYECTOS" else inv_repo.crear_semillero
    obtener = inv_repo.obtener_proyecto_por_codigo if catalogo == "PROYECTOS" else inv_repo.obtener_semillero_por_codigo
    creados = actualizados = desactivados = 0
    for r in resultados:
        if r.resultado == "CREADO":
            crear(db, r.datos["codigo"], r.datos["nombre"], r.datos["estado"])
            creados += 1
        elif r.resultado in ("ACTUALIZADO", "DESACTIVADO"):
            entidad = obtener(db, r.datos["codigo"])
            entidad.nombre = r.datos["nombre"]
            entidad.estado = r.datos["estado"]
            if r.resultado == "DESACTIVADO":
                desactivados += 1
            else:
                actualizados += 1
    return creados, actualizados, desactivados


def _confirmar_equipos(db, resultados) -> tuple[int, int]:
    creados = actualizados = 0
    for r in resultados:
        datos = r.datos
        fecha_compra = date.fromisoformat(datos["fecha_compra"]) if datos["fecha_compra"] else None
        if r.resultado == "CREADO":
            recurso = rec_repo.crear_recurso(db, datos["id_unidad"], "EQUIPO")
            rec_repo.crear_equipo(
                db, recurso.id, nombre_equipo=datos["nombre_equipo"], placa=datos["placa"],
                bodega=datos["bodega"], centro_costo=datos["centro_costo"], fecha_compra=fecha_compra,
            )
            creados += 1
        elif r.resultado == "ACTUALIZADO":
            equipo = rec_repo.obtener_equipo_por_placa(db, datos["placa"])
            equipo.nombre_equipo = datos["nombre_equipo"]
            equipo.bodega = datos["bodega"]
            equipo.centro_costo = datos["centro_costo"]
            equipo.fecha_compra = fecha_compra
            actualizados += 1
    return creados, actualizados


def confirmar_importacion(db, id_importacion: int, contexto: ContextoAutenticado) -> dict:
    importacion = repo.obtener_importacion(db, id_importacion)
    if importacion is None:
        raise NoEncontrado()
    if importacion.confirmado_at is not None:
        raise Conflicto("La importación ya fue confirmada.")

    resultados = repo.resultados_de_importacion(db, id_importacion)
    if any(r.resultado == "ERROR" for r in resultados):
        raise Conflicto("La carga contiene filas en error y no puede confirmarse.")

    if importacion.catalogo == "EQUIPOS":
        creados, actualizados = _confirmar_equipos(db, resultados)
        desactivados = 0
    else:
        creados, actualizados, desactivados = _confirmar_investigacion(db, importacion.catalogo, resultados)

    importacion.registros_creados = creados
    importacion.registros_actualizados = actualizados
    importacion.registros_desactivados = desactivados
    importacion.confirmado_at = _ahora()
    db.commit()
    db.refresh(importacion)
    return _resumen(importacion)


# --- §4.3 Historial ------------------------------------------------------------------


def _resumen(importacion) -> dict:
    return {
        "id": importacion.id, "actor_cuenta_id": importacion.actor_cuenta_id, "catalogo": importacion.catalogo,
        "archivo_referencia": importacion.archivo_referencia, "registros_creados": importacion.registros_creados,
        "registros_actualizados": importacion.registros_actualizados, "registros_desactivados": importacion.registros_desactivados,
        "created_at": importacion.created_at, "confirmado_at": importacion.confirmado_at,
    }


def detalle_importacion(db, id_importacion: int) -> dict:
    importacion = repo.obtener_importacion(db, id_importacion)
    if importacion is None:
        raise NoEncontrado()
    resultados = repo.resultados_de_importacion(db, id_importacion)
    return {
        **_resumen(importacion),
        "resultados": [{"numero_fila": r.numero_fila, "codigo": r.codigo, "resultado": r.resultado, "detalle": r.detalle} for r in resultados],
    }


def listar_importaciones(db, *, catalogo: str | None, desde: datetime | None, hasta: datetime | None, paginacion) -> tuple[list[dict], int]:
    filas, total = repo.listar_importaciones(
        db, catalogo=catalogo, desde=desde, hasta=hasta,
        limite=paginacion.tamano, desplazamiento=paginacion.offset,
    )
    return [_resumen(f) for f in filas], total
