"""Exportación del resumen del dashboard admin/gestor a CSV y Excel.

Reutiliza `AdminDashboardSummary` (misma fuente de datos que ya ve el
usuario en pantalla vía `_construir_resumen`) -- no hay lógica de
agregación nueva acá, solo serialización a dos formatos de archivo.
"""

import csv
import io

from openpyxl import Workbook

from app.schemas.admin_dashboard import AdminDashboardSummary

_ENCABEZADO_RESUMEN = ("Métrica", "Valor")


def _filas_resumen(resumen: AdminDashboardSummary) -> list[tuple[str, object]]:
    filas = [
        ("Total de reservas", resumen.total_reservas),
        ("Reservas pendientes", resumen.reservas_pendientes),
        ("Recursos activos", resumen.recursos_activos),
        ("Usuarios", resumen.usuarios),
    ]
    if resumen.laboratorio_nombre is not None:
        filas.insert(0, ("Laboratorio", resumen.laboratorio_nombre))
    return filas


def construir_csv(resumen: AdminDashboardSummary) -> bytes:
    buffer = io.StringIO()
    writer = csv.writer(buffer)

    writer.writerow(_ENCABEZADO_RESUMEN)
    writer.writerows(_filas_resumen(resumen))

    writer.writerow([])
    writer.writerow(("Reservas por estado",))
    writer.writerow(("Pendientes", "Aprobadas", "Rechazadas", "Canceladas"))
    writer.writerow(
        (
            resumen.reservas_por_estado.pendientes,
            resumen.reservas_por_estado.aprobadas,
            resumen.reservas_por_estado.rechazadas,
            resumen.reservas_por_estado.canceladas,
        )
    )

    writer.writerow([])
    writer.writerow(("Reservas por fecha",))
    writer.writerow(("Fecha", "Cantidad"))
    for item in resumen.reservas_por_fecha:
        writer.writerow((item.fecha.isoformat(), item.cantidad))

    writer.writerow([])
    writer.writerow(("Reservas por laboratorio",))
    writer.writerow(("Laboratorio", "Cantidad"))
    for item in resumen.reservas_por_laboratorio:
        writer.writerow((item.nombre, item.cantidad))

    writer.writerow([])
    writer.writerow(("Recursos más reservados",))
    writer.writerow(("Recurso", "Cantidad"))
    for item in resumen.recursos_mas_reservados:
        writer.writerow((item.nombre, item.cantidad))

    writer.writerow([])
    writer.writerow(("Ocupación por día y hora",))
    writer.writerow(("Día", "Hora", "Cantidad"))
    for item in resumen.ocupacion_por_dia_hora:
        writer.writerow((item.dia, item.hora, item.cantidad))

    return buffer.getvalue().encode("utf-8-sig")


def construir_xlsx(resumen: AdminDashboardSummary) -> bytes:
    libro = Workbook()

    hoja_resumen = libro.active
    hoja_resumen.title = "Resumen"
    hoja_resumen.append(_ENCABEZADO_RESUMEN)
    for fila in _filas_resumen(resumen):
        hoja_resumen.append(fila)

    hoja_estado = libro.create_sheet("Reservas por estado")
    hoja_estado.append(("Pendientes", "Aprobadas", "Rechazadas", "Canceladas"))
    hoja_estado.append(
        (
            resumen.reservas_por_estado.pendientes,
            resumen.reservas_por_estado.aprobadas,
            resumen.reservas_por_estado.rechazadas,
            resumen.reservas_por_estado.canceladas,
        )
    )

    hoja_fecha = libro.create_sheet("Reservas por fecha")
    hoja_fecha.append(("Fecha", "Cantidad"))
    for item in resumen.reservas_por_fecha:
        hoja_fecha.append((item.fecha.isoformat(), item.cantidad))

    hoja_laboratorio = libro.create_sheet("Reservas por laboratorio")
    hoja_laboratorio.append(("Laboratorio", "Cantidad"))
    for item in resumen.reservas_por_laboratorio:
        hoja_laboratorio.append((item.nombre, item.cantidad))

    hoja_recursos = libro.create_sheet("Recursos más reservados")
    hoja_recursos.append(("Recurso", "Cantidad"))
    for item in resumen.recursos_mas_reservados:
        hoja_recursos.append((item.nombre, item.cantidad))

    hoja_ocupacion = libro.create_sheet("Ocupación por día y hora")
    hoja_ocupacion.append(("Día", "Hora", "Cantidad"))
    for item in resumen.ocupacion_por_dia_hora:
        hoja_ocupacion.append((item.dia, item.hora, item.cantidad))

    buffer = io.BytesIO()
    libro.save(buffer)
    return buffer.getvalue()
